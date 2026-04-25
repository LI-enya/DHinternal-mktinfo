"""
brands.json から各ブランドの「竞争对手列表」を出力するスクリプト。

出力:
  1. Excel: _品牌资料库/IO-ISOLE/竞争对手列表.xlsx
  2. Excel: _品牌资料库/LA-Laise/竞争对手列表.xlsx
  3. 飞书テキスト: stdout（コピー＆ペースト用）

Usage:
  python scripts/export_brands.py
"""
import json
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

REPO = Path(__file__).resolve().parent.parent
BRANDS_PATH = REPO / "brands.json"

WORKSPACE = REPO.parent.parent  # AI/
LIBRARY = WORKSPACE / "_品牌资料库"

# 子分类の日本語ラベル
TIER_LABELS = {
    "core": "最核心",
    "price-aligned": "価格段対",
    "tone-aligned": "調性対",
    "trend-watch": "動向参考",
    "activity-watch": "活動観察",
}

# 描述用の中文标签（飞书文本用）
TIER_LABELS_ZH = {
    "core": "最核心",
    "price-aligned": "价格段对",
    "tone-aligned": "调性对",
    "trend-watch": "动向参考",
    "activity-watch": "活动观察",
}


def collect_rows(brands, category_key, brand_label):
    """brands.json の bag or shoe セクションを Excel 行リストに変換"""
    section = brands[category_key]
    rows = []

    # Primary (CEO提供)
    for entry in section["primary"]:
        rows.append({
            "name": entry["name"],
            "url": entry.get("officialUrl", ""),
            "source": "CEO提供",
            "category": brand_label,
            "tier": "Primary",
            "subTier": TIER_LABELS_ZH.get(entry["tier"], entry["tier"]),
            "reason": entry.get("reason", ""),
        })

    # Reference seed (CEO提供)
    for entry in section["reference"]["seed"]:
        rows.append({
            "name": entry["name"],
            "url": entry.get("officialUrl", ""),
            "source": "CEO提供",
            "category": brand_label,
            "tier": "Reference",
            "subTier": TIER_LABELS_ZH.get(entry["tier"], entry["tier"]),
            "reason": entry.get("reason", ""),
        })

    # Reference aiSuggested (Agent建议)
    for entry in section["reference"].get("aiSuggested", []):
        rows.append({
            "name": entry["name"],
            "url": entry.get("officialUrl", ""),
            "source": "AI建议",
            "category": brand_label,
            "tier": "Reference",
            "subTier": TIER_LABELS_ZH.get(entry["tier"], entry["tier"]),
            "reason": entry.get("reason", ""),
        })

    return rows


def write_excel(rows, output_path, brand_name):
    """Excel ファイルを書き出す"""
    wb = Workbook()
    ws = wb.active
    ws.title = "竞争对手列表"

    headers = ["品牌名", "URL", "来源", "类别", "层级", "子分类", "选取理由"]
    ws.append(headers)

    # ヘッダーフォーマット
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", fgColor="4A6FA5")
    header_align = Alignment(horizontal="center", vertical="center")
    thin = Side(border_style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = border

    # 行データ
    ai_fill = PatternFill("solid", fgColor="FFF8E7")  # AI建议は淡い黄色
    primary_fill = PatternFill("solid", fgColor="FDF4F0")  # Primary は淡いオレンジ

    for row in rows:
        ws.append([
            row["name"],
            row["url"],
            row["source"],
            row["category"],
            row["tier"],
            row["subTier"],
            row["reason"],
        ])
        row_idx = ws.max_row
        # AI建议行を着色
        if row["source"] == "AI建议":
            for col in range(1, len(headers) + 1):
                ws.cell(row=row_idx, column=col).fill = ai_fill
        elif row["tier"] == "Primary":
            for col in range(1, len(headers) + 1):
                ws.cell(row=row_idx, column=col).fill = primary_fill
        # 罫線
        for col in range(1, len(headers) + 1):
            ws.cell(row=row_idx, column=col).border = border
            ws.cell(row=row_idx, column=col).alignment = Alignment(vertical="top", wrap_text=True)

    # 列幅調整
    widths = [22, 50, 12, 8, 12, 12, 60]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + i)].width = w

    # ヘッダー行のフリーズ
    ws.freeze_panes = "A2"

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    print(f"[OK] {output_path.name} に {len(rows)} 行書き込み: {output_path}")


def render_feishu_text(brands, category_key, brand_display):
    """飞书メッセージ用テキストを生成（CEO先 → AI後の順）"""
    section = brands[category_key]
    primary = section["primary"]
    ref_seed = section["reference"]["seed"]
    ref_ai = section["reference"].get("aiSuggested", [])
    total = len(primary) + len(ref_seed) + len(ref_ai)

    lines = []
    lines.append(f"🎯 {brand_display}（共{total}個）")
    lines.append("━" * 30)
    lines.append("")

    # ━━ Primary ━━
    lines.append(f"━━ Primary（{len(primary)}）━━")
    by_subtier = {}
    for entry in primary:
        by_subtier.setdefault(entry["tier"], []).append(entry)
    # 子分類順序: core → price-aligned
    order = ["core", "price-aligned"]
    for subtier in order:
        if subtier not in by_subtier:
            continue
        lines.append("")
        lines.append(f"【{TIER_LABELS_ZH[subtier]}】")
        for entry in by_subtier[subtier]:
            lines.append(f"• {entry['name']}; {entry.get('officialUrl','')}")

    lines.append("")
    # ━━ Reference ━━
    ref_total = len(ref_seed) + len(ref_ai)
    lines.append(f"━━ Reference（{ref_total}）━━")
    # CEO + AI を子分類別にグループ化、CEO 先・AI 後
    ceo_by_subtier = {}
    ai_by_subtier = {}
    for entry in ref_seed:
        ceo_by_subtier.setdefault(entry["tier"], []).append(entry)
    for entry in ref_ai:
        ai_by_subtier.setdefault(entry["tier"], []).append(entry)

    # 子分類順序: tone-aligned → trend-watch → activity-watch
    ref_order = ["tone-aligned", "trend-watch", "activity-watch"]
    for subtier in ref_order:
        if subtier not in ceo_by_subtier and subtier not in ai_by_subtier:
            continue
        lines.append("")
        lines.append(f"【{TIER_LABELS_ZH[subtier]}】")
        # CEO 先
        for entry in ceo_by_subtier.get(subtier, []):
            lines.append(f"• {entry['name']}; {entry.get('officialUrl','')}")
        # AI 後（括号で AI建议マーク）
        for entry in ai_by_subtier.get(subtier, []):
            lines.append(f"• {entry['name']}（AI建议）; {entry.get('officialUrl','')}")

    return "\n".join(lines)


def main():
    brands = json.loads(BRANDS_PATH.read_text(encoding="utf-8"))

    # === IO (ISOLÉ / bag) ===
    io_rows = collect_rows(brands, "bag", "包")
    io_excel = LIBRARY / "IO-ISOLE" / "竞争对手列表.xlsx"
    write_excel(io_rows, io_excel, "ISOLÉ")
    io_text = render_feishu_text(brands, "bag", "ISOLÉ（バッグ・¥30K前後）")

    # === LA (l'aisé / shoe) ===
    la_rows = collect_rows(brands, "shoe", "鞋")
    la_excel = LIBRARY / "LA-Laise" / "竞争对手列表.xlsx"
    write_excel(la_rows, la_excel, "l'aisé")
    la_text = render_feishu_text(brands, "shoe", "l'aisé（シューズ・¥20K前後・dress shoes強化）")

    # 飞书テキストをファイルにも保存（バックアップ）
    (REPO / "scripts" / "feishu_io.txt").write_text(io_text, encoding="utf-8")
    (REPO / "scripts" / "feishu_la.txt").write_text(la_text, encoding="utf-8")

    print()
    print("=" * 60)
    print("【飞书文本】 ISOLÉ")
    print("=" * 60)
    print(io_text)
    print()
    print("=" * 60)
    print("【飞书文本】 l'aisé")
    print("=" * 60)
    print(la_text)


if __name__ == "__main__":
    main()
